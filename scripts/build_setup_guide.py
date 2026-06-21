#!/usr/bin/env python3

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "out" / "retail" / "Pixel-Meadow-Setup-Guide.docx"
DAY_IMAGE = ROOT / "public" / "wallpapers" / "day" / "source.png"
NIGHT_IMAGE = ROOT / "public" / "wallpapers" / "night" / "source.png"

NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
GREEN = RGBColor(43, 122, 75)
MUTED = RGBColor(95, 105, 115)
LIGHT_BLUE = "E8EEF5"


def set_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_margins(cell, top=120, start=160, bottom=120, end=160):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_body(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_font(first, size=11, color=NAVY, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_font(rest, size=11, color=NAVY)
    else:
        run = paragraph.add_run(text)
        set_font(run, size=11, color=NAVY)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.375)
    paragraph.paragraph_format.first_line_indent = Inches(-0.188)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    set_font(run, size=11, color=NAVY)


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.left_indent = Inches(0.375)
    paragraph.paragraph_format.first_line_indent = Inches(-0.188)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    set_font(run, size=11, color=NAVY)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    sizes = {1: 16, 2: 13, 3: 12}
    colors = {1: BLUE, 2: BLUE, 3: NAVY}
    set_font(run, size=sizes[level], color=colors[level], bold=True)
    return paragraph


def add_note(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    table.rows[0].cells[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    shade_cell(cell, LIGHT_BLUE)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    label = paragraph.add_run(f"{title}: ")
    set_font(label, size=10.5, color=GREEN, bold=True)
    body = paragraph.add_run(text)
    set_font(body, size=10.5, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = NAVY
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    spacing = {1: (18, 10), 2: (14, 7), 3: (10, 5)}
    for level in (1, 2, 3):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.size = Pt({1: 16, 2: 13, 3: 12}[level])
        style.font.color.rgb = {1: BLUE, 2: BLUE, 3: NAVY}[level]
        style.font.bold = True
        style.paragraph_format.space_before = Pt(spacing[level][0])
        style.paragraph_format.space_after = Pt(spacing[level][1])

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("QUIETPIXELSTUDIO  /  PIXEL MEADOW")
    set_font(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    run = footer.add_run("Pixel Meadow Setup Guide  |  v1.0  |  2026")
    set_font(run, size=8.5, color=MUTED)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(2)
    run = kicker.add_run("CUSTOMER SETUP GUIDE")
    set_font(run, size=10, color=GREEN, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("Pixel Meadow")
    set_font(run, size=30, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("Automatic day + night animated wallpapers for Mac")
    set_font(run, size=13.5, color=MUTED)

    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_after = Pt(12)
    picture.add_run().add_picture(str(DAY_IMAGE), width=Inches(6.1))

    add_note(doc, "Designed for", "Apple Silicon Macs running macOS 15 or later with the free Aerial 4 app.")

    add_heading(doc, "Before you begin", 1)
    for item in (
        "Apple Silicon Mac (M1 or newer) with macOS 15 or later",
        "Free Aerial 4 app from aerialscreensaver.github.io",
        "Both Pixel Meadow MP4 files",
        "The unzipped Pixel Meadow Automation folder",
    ):
        add_bullet(doc, item)
    add_body(doc, "Aerial is a separate open-source project. QuietPixelStudio is not affiliated with or endorsed by Aerial.")

    add_heading(doc, "Install Pixel Meadow", 1)
    for step in (
        "Install and open Aerial once.",
        "Put both MP4 files inside the unzipped Pixel Meadow Automation folder.",
        "Control-click Install Pixel Meadow.command, choose Open, and confirm the macOS prompt. The installer does not require an administrator password.",
        "In Aerial, select My Videos, set playback speed to 1.0x, mute audio, and start Wallpaper mode.",
    ):
        add_number(doc, step)

    add_note(doc, "How it works", "Every five minutes, the automation checks sunrise and sunset, keeps one scene active for seamless looping, and switches Aerial only when the scene actually changes.")

    add_heading(doc, "Recommended Aerial settings", 1)
    for item in (
        "Filter: My Videos",
        "Playback speed: 1.0x",
        "Audio: muted",
        "Wallpaper: enabled",
        "Pause on battery or fullscreen: your preference",
    ):
        add_bullet(doc, item)

    add_heading(doc, "If macOS blocks the installer", 1)
    add_body(doc, "The installer is an unsigned shell script, not a notarized Mac app. Control-click the file and choose Open. If macOS still blocks it, open System Settings > Privacy & Security and choose Open Anyway for the Pixel Meadow installer.")
    add_note(doc, "Safety", "Only open the installer from your original Etsy download. Never bypass a macOS warning for a copy received elsewhere.")

    add_heading(doc, "Troubleshooting", 1)
    issues = (
        ("The wallpaper moves too slowly", "Set Aerial playback speed to 1.0x. The slower motion is already encoded into the video."),
        ("The screen flashes black every 18.5 seconds", "Only one Pixel Meadow file should be in /Users/Shared/Aerial/My Videos/. Run the installer again to restore the single-active-file configuration."),
        ("The wrong scene is active", "Confirm macOS date and time are automatic, open Aerial once, and wait up to five minutes. The fallback schedule is 7:00 AM to 7:00 PM."),
        ("Aerial reports no available video", "Open Aerial's Video Library and select My Videos. If neither file appears, run the installer again with both MP4 files beside it."),
    )
    for title_text, detail in issues:
        add_heading(doc, title_text, 2)
        add_body(doc, detail)

    add_heading(doc, "Remove the automation", 1)
    add_body(doc, "Control-click Uninstall Pixel Meadow.command and choose Open. The automation and LaunchAgent are removed. Both MP4 files are preserved in Movies/Pixel Meadow Wallpapers.")

    add_heading(doc, "Support", 1)
    add_body(doc, "When requesting help, include your Mac model, macOS version, Aerial version, and a screenshot of the issue. Do not send passwords or private system data.")

    closing = doc.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    closing.paragraph_format.space_before = Pt(12)
    closing.paragraph_format.space_after = Pt(0)
    closing.add_run().add_picture(str(NIGHT_IMAGE), width=Inches(6.1))

    doc.core_properties.title = "Pixel Meadow Setup Guide"
    doc.core_properties.subject = "QuietPixelStudio customer installation guide"
    doc.core_properties.author = "QuietPixelStudio"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
